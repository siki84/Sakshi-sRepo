"""Markdown <-> .docx, so drafts arrive as editable Word files you can change and send back."""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

BOLD = re.compile(r"\*\*(.+?)\*\*")


def _add_runs(paragraph, text: str) -> None:
    pos = 0
    for m in BOLD.finditer(text):
        if m.start() > pos:
            paragraph.add_run(text[pos : m.start()])
        paragraph.add_run(m.group(1)).bold = True
        pos = m.end()
    if pos < len(text):
        paragraph.add_run(text[pos:])


def markdown_to_docx(markdown: str, path: Path) -> Path:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)
    first_heading = True

    for raw in markdown.splitlines():
        line = raw.rstrip()
        if not line.strip():
            continue
        if line.startswith("# "):
            p = doc.add_heading(line[2:].strip(), level=0 if first_heading else 1)
            if first_heading:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            first_heading = False
        elif line.startswith("## "):
            doc.add_heading(line[3:].strip(), level=1)
        elif line.startswith("### "):
            doc.add_heading(line[4:].strip(), level=2)
        elif re.match(r"^\s*[-*] ", line):
            _add_runs(doc.add_paragraph(style="List Bullet"), re.sub(r"^\s*[-*] ", "", line))
        else:
            _add_runs(doc.add_paragraph(), line)

    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)
    return path


def docx_to_text(path: Path) -> str:
    """Plain text of a .docx (used to read your edited versions back in)."""
    doc = Document(path)
    lines = []
    for p in doc.paragraphs:
        text = p.text.strip()
        if not text:
            continue
        style = (p.style.name or "").lower()
        if style == "title" or style == "heading 1" and not lines:
            lines.append(f"# {text}")
        elif style == "heading 1":
            lines.append(f"## {text}")
        elif style.startswith("heading"):
            lines.append(f"### {text}")
        elif "list" in style:
            lines.append(f"- {text}")
        else:
            lines.append(text)
    return "\n".join(lines)
